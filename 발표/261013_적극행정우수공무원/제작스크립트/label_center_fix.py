from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
X=lambda e:e*1920/12192000
p=Presentation('fixed.pptx');s=p.slides[3]
cards=[q for q in s.shapes if q.shape_type==1 and 'Rounded' in q.name and abs(X(q.top)-290)<2 and X(q.width)>500]
for c in cards:
    cx=c.left+c.width/2
    for q in s.shapes:
        if q.has_text_frame and q is not c and abs(X(q.top+q.height/2)-335)<20 and c.left<=q.left+q.width/2<=c.left+c.width:
            q.left,q.top,q.width,q.height=c.left,c.top,c.width,c.height
            tf=q.text_frame;tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            for m in ('margin_left','margin_right','margin_top','margin_bottom'):setattr(tf,m,0)
            for para in tf.paragraphs:
                para.alignment=PP_ALIGN.CENTER;para.line_spacing=None
                for r in para.runs: print(repr(r.text),r.font.size.pt if r.font.size else None); r.font.size=Pt(25)
p.save('fixed2.pptx')
