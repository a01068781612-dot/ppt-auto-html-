import json,sys,os,re
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
D,JS,OUT=sys.argv[1:4]
data=json.load(open(JS));os.makedirs(D+'/ed/img',exist_ok=True)
E=lambda v:Emu(int(round(v*12192000/1920)))
prs=Presentation();prs.slide_width=Emu(12192000);prs.slide_height=Emu(6858000)
def fam(w):
    if w>=850:return 'Pretendard Black',False
    if w>=750:return 'Pretendard ExtraBold',False
    if w>=650:return 'Pretendard',True
    if w>=550:return 'Pretendard SemiBold',False
    if w>=450:return 'Pretendard Medium',False
    return 'Pretendard',False
def setfont(run,r):
    f=run.font;name,b=fam(r['wt']);f.size=Pt(r['size']*0.5);f.bold=b;f.color.rgb=RGBColor.from_string(r['c'])
    rPr=run._r.get_or_add_rPr()
    for tag in ('latin','ea','cs'):
        el=rPr.find(qn('a:'+tag))
        if el is None:el=rPr.makeelement(qn('a:'+tag),{});rPr.append(el)
        el.set('typeface',name)
    if r['ls']:rPr.set('spc',str(int(round(r['ls']*0.5*100))))
    if r.get('ghost'):
        for el in rPr.findall(qn('a:solidFill')):rPr.remove(el)
        ln=rPr.makeelement(qn('a:ln'),{'w':'12700'});sf=ln.makeelement(qn('a:solidFill'),{});c=sf.makeelement(qn('a:srgbClr'),{'val':r['c']});sf.append(c);ln.append(sf)
        nf=rPr.makeelement(qn('a:noFill'),{});rPr.insert(0,nf);rPr.insert(0,ln)
def img(it,n):
    src=D+'/'+it['src'];im=Image.open(src).convert('RGB');W,H=im.size;bw,bh=it['w'],it['h']
    if it['fit']=='cover':
        s=max(bw/W,bh/H);cw,ch=bw/s,bh/s;pos=it['pos'].split()
        def off(p,free):
            if p.endswith('%'):return free*float(p[:-1])/100
            if p.endswith('px'):return float(p[:-2])/s
            return free/2
        ox=off(pos[0],W-cw);oy=off(pos[1] if len(pos)>1 else '50%',H-ch)
        im=im.crop((int(ox),int(oy),int(ox+cw),int(oy+ch)))
    sc=2;im=im.resize((max(1,int(bw*sc)),max(1,int(bh*sc))),Image.LANCZOS)
    p=f"{D}/ed/img/{n}"
    if it['rad']>0:
        m=Image.new('L',im.size,0);ImageDraw.Draw(m).rounded_rectangle((0,0,im.size[0]-1,im.size[1]-1),radius=int(it['rad']*sc),fill=255)
        im=im.convert('RGBA');im.putalpha(m);p+='.png';im.save(p)
    else:p+='.jpg';im.save(p,quality=90)
    return p
SH={'chevron':MSO_SHAPE.CHEVRON,'homePlate':MSO_SHAPE.PENTAGON,'notchL':MSO_SHAPE.CHEVRON}
notes={'v1':'[영상1 · 45초] 현 문제 실태 — 삽입 → 비디오 → 이 디바이스, 화면 전체 크기, 재생: 자동 실행.',
'v2':'[영상2 · 1분 20초] 내비 실제 구동 → 농골길 실제 구동 → 원패스 설계 시연 → 엔딩 "그 1분이, 누군가에게는 평생입니다". 삽입 → 비디오, 화면 전체 크기, 자동 실행.',
's3':'점선 칸 4개 = 필요 자료 자리. 칸을 지우고 같은 자리에 삽입 → 그림.'}
for sid,sd in data.items():
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(sd['bg'])
    for i,it in enumerate(sd['items']):
        k=it['k']
        if k=='rect':
            sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,E(it['x']),E(it['y']),E(it['w']),E(it['h']))
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(it['fill']);sh.line.fill.background();sh.shadow.inherit=False
        elif k=='shape':
            if it['shape']:t=SH[it['shape']]
            elif it['ellipse']:t=MSO_SHAPE.OVAL
            elif it['rad']>0:t=MSO_SHAPE.ROUNDED_RECTANGLE
            else:t=MSO_SHAPE.RECTANGLE
            sh=s.shapes.add_shape(t,E(it['x']),E(it['y']),E(it['w']),E(it['h']));sh.shadow.inherit=False
            if t==MSO_SHAPE.ROUNDED_RECTANGLE:sh.adjustments[0]=min(0.5,it['rad']/min(it['w'],it['h']))
            if t in(MSO_SHAPE.CHEVRON,MSO_SHAPE.PENTAGON):sh.adjustments[0]=0.06*it['w']/it['h']
            if it['fill']:sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(it['fill'])
            else:sh.fill.background()
            if it['line']:
                sh.line.color.rgb=RGBColor.from_string(it['line']);sh.line.width=Pt(max(0.5,it['lw']*0.5))
                if it['dash']:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
            else:sh.line.fill.background()
        elif k=='img':
            s.shapes.add_picture(img(it,f'{sid}_{i}'),E(it['x']),E(it['y']),E(it['w']),E(it['h']))
        elif k=='text':
            runs=it['runs'];lines=[[]]
            for r in runs:
                if r.get('br'):lines.append([])
                else:lines[-1].append(dict(r))
            for ln in lines:
                if ln:ln[0]['t']=ln[0]['t'].lstrip();ln[-1]['t']=ln[-1]['t'].rstrip()
            lines=[[r for r in ln if r['t']] for ln in lines]
            while lines and not lines[-1]:lines.pop()
            while lines and not lines[0]:lines.pop(0)
            if not lines:continue
            wrap=it['lines']>it['brs']+1
            x,w=it['x'],it['w']
            if wrap:w=w*1.04+4
            else:
                ex=w*0.15+10
                if it['align']=='center':x-=ex/2
                elif it['align']=='right':x-=ex
                w+=ex
            tb=s.shapes.add_textbox(E(x),E(it['y']),E(w),E(it['h']));tf=tb.text_frame
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=wrap;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.TOP
            lh=it['lh'];multi=len(lines)>1 or wrap
            for j,ln in enumerate(lines):
                para=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                para.alignment={'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(it['align'],PP_ALIGN.LEFT)
                if multi and lh.endswith('px'):para.line_spacing=Pt(float(lh[:-2])*0.5)
                for r in ln:
                    run=para.add_run();run.text=r['t'];setfont(run,r)
    for shp in s.shapes:
        st=shp._element.find(qn('p:style'))
        if st is not None:shp._element.remove(st)
    if sid in notes:s.notes_slide.notes_text_frame.text=notes[sid]
prs.save(OUT);print('saved')
