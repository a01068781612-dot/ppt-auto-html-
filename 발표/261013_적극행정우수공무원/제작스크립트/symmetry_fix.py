from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.enum.text import PP_ALIGN
P=lambda v:Emu(int(round(v*12192000/1920)))
X=lambda e:e*1920/12192000
p=Presentation('u.pptx')
def geo(sh):return X(sh.left),X(sh.top),X(sh.width),X(sh.height)
def setg(sh,x,y,w,h):sh.left,sh.top,sh.width,sh.height=P(x),P(y),P(w),P(h)
def fitpic(pic,x,y,w,h,anchor='c'):
    x0,y0,w0,h0=geo(pic);a0=w0/h0;a1=w/h
    cl,ct,cr,cb=pic.crop_left,pic.crop_top,pic.crop_right,pic.crop_bottom
    fw=1-cl-cr;fh=1-ct-cb
    if a0>a1:  # trim horizontally
        keep=a1/a0;t=fw*(1-keep);cl+=t/2;cr+=t/2;trim=('x',(1-keep)/2,0)
    else:
        keep=a0/a1;t=fh*(1-keep)
        if anchor=='top_keep_bottom': ct+=t;trim=('y',1-keep,0)
        else: ct+=t/2;cb+=t/2;trim=('y',(1-keep)/2,0)
    pic.crop_left,pic.crop_top,pic.crop_right,pic.crop_bottom=cl,ct,cr,cb
    setg(pic,x,y,w,h);return (x0,y0,w0,h0),keep,trim
def remove(sh):sh._element.getparent().remove(sh._element)
# ---------- slide 3: center dot grid in card 2
s=p.slides[2]
for sh in s.shapes:
    if sh.name.startswith('Oval'):
        sh.left=sh.left+P(3)
# ---------- slide 4
s=p.slides[3];sh=list(s.shapes)
rows={'lab':(290,90),'r2':(410,275),'r3':(715,275)}
cols=[110,690,1270];CW=540
def colof(x):return min(range(3),key=lambda i:abs(x-(cols[i]+CW/2)))
# remove shapes hidden under pictures
pics=[q for q in sh if q.shape_type==13]
def covered(q):
    x,y,w,h=geo(q)
    for pc in pics:
        a,b,c,d=geo(pc)
        if q is not pc and x>=a-2 and y>=b-2 and x+w<=a+c+2 and y+h<=b+d+2: return True
    return False
hidden=[q for q in sh if q.shape_type!=13 and geo(q)[1]>440 and covered(q)]
for q in hidden: remove(q)
sh=list(s.shapes)
def ymap(y):
    bp=[(290,290),(400,380),(448,410),(728,685),(776,715),(1056,990)]
    for (a,A),(b,B) in zip(bp,bp[1:]):
        if a<=y<=b:return A+(y-a)*(B-A)/(b-a)
    return y
hl=None
for q in sh:
    x,y,w,h=geo(q)
    if y<280 or q.name.startswith('TextBox 64'):continue
    c=colof(x+w/2)
    if q.shape_type==1 and 'Rounded' in q.name and w>500:   # cards
        if h<150: setg(q,cols[c],290,CW,90)
        elif y<600: setg(q,cols[c],410,CW,275)
        else: setg(q,cols[c],715,CW,275)
    elif q.shape_type==13:
        top=410 if y<600 else 715
        if q.name=='Picture 18':
            old,keep,trim=fitpic(q,cols[c]+10,top+10,CW-20,255,'top_keep_bottom');pic18=(old,keep)
        else: fitpic(q,cols[c]+10,top+10,CW-20,255)
    elif q.shape_type==1 and w>500 and h<80: hl=q
    elif q.has_text_frame and q.text_frame.text.strip() in ('▲','+'):
        cy=395 if q.text_frame.text.strip()=='▲' else 700
        setg(q,cols[c]+CW/2-20,cy-18,40,36)
        for para in q.text_frame.paragraphs:
            para.alignment=PP_ALIGN.CENTER
    elif q.has_text_frame and 300<y<400:   # card labels -> centered
        setg(q,cols[c],290+45-h/2,CW,h)
        q.text_frame.word_wrap=False
        for para in q.text_frame.paragraphs: para.alignment=PP_ALIGN.CENTER
    else:
        nc=ymap(y+h/2);q.top=P(nc-h/2)
        if q.shape_type==1 and w>200 and h>200: # placeholder box
            setg(q,x,715+16,w,275-32)
# red highlight follows the nav picture
if hl is not None:
    (x0,y0,w0,h0),keep=pic18
    x,y,w,h=geo(hl);nx,ny,nw,nh=cols[0]+10,420,520,255
    sx=nw/w0;sy=nh/(h0*keep);vis_top=y0+h0*(1-keep)
    setg(hl,nx+(x-x0)*sx,ny+(y-vis_top)*sy,w*sx,h*sy)
# ---------- slide 7
s=p.slides[6];sh=list(s.shapes)
PX=[110,687,1263];PW=547;PY,PH=290,390
frames=[q for q in sh if q.shape_type==1 and 'Rounded' in q.name and 490<X(q.width)<520]
frames.sort(key=lambda q:q.left)
for i,f in enumerate(frames): setg(f,PX[i],PY,PW,PH)
ps=sorted([q for q in sh if q.shape_type==13],key=lambda q:q.left)
for i,q in enumerate(ps): fitpic(q,PX[i]+1,PY+1,PW-2,PH-2)
for q in sh:
    if not q.has_text_frame:continue
    x,y,w,h=geo(q);t=q.text_frame.text
    if 600<y<700:
        i=min(range(3),key=lambda k:abs(x-[248,778,1308][k]))
        q.left=P(PX[i]);q.top=P(PY+PH+(16 if y<640 else 49))
# pill -> horizontal chip on first photo
pill=[q for q in sh if q.shape_type==1 and 'Rounded' in q.name and X(q.width)<130][0]
pt=[q for q in sh if q.has_text_frame and q.text_frame.text.replace('\n','')=='경남에서증명'][0]
setg(pill,PX[0]+18,PY+18,200,52)
tf=pt.text_frame;r0=tf.paragraphs[0].runs[0]
for para in tf.paragraphs[1:]: para._p.getparent().remove(para._p)
for r in tf.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
r0.text='경남에서 증명';tf.word_wrap=False;tf.paragraphs[0].alignment=PP_ALIGN.CENTER
setg(pt,PX[0]+18,PY+18+(52-30)/2,200,30)
# keep pill above picture: move to end of spTree
for el in (pill._element,pt._element):
    el.getparent().remove(el);s.shapes._spTree.append(el)
p.save('fixed.pptx');print('saved, removed hidden on 4p:',len(hidden))
