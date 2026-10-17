# v.2.3.0 → v.2.4.0: 4장 태그 위치 + 전체 모션(입장 애니메이션 자동 재생 + 화면 전환)
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

SRC, OUT = sys.argv[1], sys.argv[2]
NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
p = Presentation(SRC)
S = lambda n: p.slides[n - 1]
P = lambda v: Emu(int(round(v * 9525)))

def sh(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid: return s
    raise KeyError(sid)

# ---------- 4장: 태그 위치(사용자 시안) ----------
s4 = S(4)
t = sh(s4, 92); t.left, t.top = P(1050), P(493)      # 기존 · 정부24 원스톱 → 오른쪽
t = sh(s4, 89); t.left, t.top = P(872), P(282)       # 한 층 더 · 재난 8개 기관 연결 → 가운데

# ---------- 애니메이션 XML ----------
class Tl:
    def __init__(s): s.id = 4; s.items = []; s.spids = []
    def nid(s): s.id += 1; return s.id
    def tgt(s, spid): return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    def vis(s, spid):
        return (f'<p:set><p:cBhvr><p:cTn id="{s.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'{s.tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
    def fade(s, spid, dur):
        return f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{s.nid()}" dur="{dur}"/>{s.tgt(spid)}</p:cBhvr></p:animEffect>'
    def wipe(s, spid, dur, d):
        return f'<p:animEffect transition="in" filter="wipe({d})"><p:cBhvr><p:cTn id="{s.nid()}" dur="{dur}"/>{s.tgt(spid)}</p:cBhvr></p:animEffect>'
    def anim(s, spid, dur, attr, a, b, extra=''):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{s.nid()}" dur="{dur}" fill="hold"{extra}/>{s.tgt(spid)}'
                f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                f'<p:tav tm="0"><p:val><p:strVal val="{a}"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="{b}"/></p:val></p:tav></p:tavLst></p:anim>')
    def add(s, spids, kind, t, dur=500):
        if isinstance(spids, int): spids = [spids]
        for spid in spids:
            if kind == 'fade':  pid, sub, b = 10, 0, s.fade(spid, dur)
            elif kind == 'rise': pid, sub, b = 42, 0, s.fade(spid, dur) + s.anim(spid, dur, 'ppt_y', '#ppt_y+0.06', '#ppt_y', ' decel="100000"')
            elif kind == 'zoom': pid, sub, b = 53, 16, s.anim(spid, dur, 'ppt_w', '0', '#ppt_w', ' decel="60000"') + s.anim(spid, dur, 'ppt_h', '0', '#ppt_h', ' decel="60000"') + s.fade(spid, dur)
            elif kind == 'up':   pid, sub, b = 22, 4, s.wipe(spid, dur, 'up')
            elif kind == 'right': pid, sub, b = 22, 8, s.wipe(spid, dur, 'right')
            node = 'afterEffect' if not s.items else 'withEffect'
            s.items.append(f'<p:par><p:cTn id="{s.nid()}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" grpId="0" nodeType="{node}">'
                           f'<p:stCondLst><p:cond delay="{int(t)}"/></p:stCondLst><p:childTnLst>{s.vis(spid)}{b}</p:childTnLst></p:cTn></p:par>')
            s.spids.append(spid)
    def xml(s, slide):
        sps = {e.get('id') for e in slide.element.iter(f'{{{NS}}}cNvPr') if e.getparent().getparent().tag == f'{{{NS}}}sp'}
        bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in dict.fromkeys(s.spids) if str(i) in sps)
        return (f'<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
                f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
                f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
                f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(s.items)}</p:childTnLst></p:cTn></p:par>'
                f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
                f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
                f'</p:childTnLst></p:cTn></p:par></p:tnLst>' + (f'<p:bldLst>{bld}</p:bldLst>' if bld else '') + '</p:timing>')

def apply(slide, tl):
    el = slide.element
    old = el.find(f'{{{NS}}}timing')
    if old is not None: el.remove(old)
    new = etree.fromstring(tl.xml(slide))
    ext = el.find(f'{{{NS}}}extLst')
    (ext.addprevious(new) if ext is not None else el.append(new))

def transition(slide, kind='fade'):
    el = slide.element
    old = el.find(f'{{{NS}}}transition')
    if old is not None: el.remove(old)
    MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    for ac in el.findall(f'{{{MC}}}AlternateContent'):
        if ac.find(f'.//{{{NS}}}transition') is not None: el.remove(ac)
    tr = etree.fromstring(f'<p:transition xmlns:p="{NS}" spd="med">' + ('<p:fade/>' if kind == 'fade' else '<p:push dir="u"/>') + '</p:transition>')
    anchor = el.find(f'{{{NS}}}clrMapOvr')
    if anchor is None: anchor = el.find(f'{{{NS}}}cSld')
    anchor.addnext(tr)

def header(tl):
    tl.add(2, 'fade', 0, 400); tl.add([4, 5, 6, 7], 'zoom', 0, 400)
    tl.add(8, 'rise', 150, 600); tl.add(3, 'right', 300, 900); tl.add(9, 'right', 400, 700)

# 1장 표지
tl = Tl()
tl.add(2, 'fade', 0, 400); tl.add(3, 'right', 150, 600); tl.add(4, 'rise', 300, 500)
tl.add([5, 6, 7, 8], 'zoom', 450, 400); tl.add(9, 'rise', 550, 700); tl.add(10, 'rise', 850, 700)
tl.add(14, 'zoom', 1100, 600); tl.add(15, 'zoom', 1300, 600); tl.add(16, 'zoom', 1500, 600)
tl.add([21, 22], 'fade', 1350, 400); tl.add([19, 20], 'fade', 1550, 400); tl.add([17, 18], 'fade', 1750, 400)
tl.add(11, 'fade', 1950, 500); tl.add(12, 'rise', 2150, 500); tl.add(23, 'fade', 2350, 500)
apply(S(1), tl)

# 3장 추진 배경
s3 = S(3); tl = Tl(); header(tl)
for i, (a, b) in enumerate([(10, 11), (12, 13), (14, 15)]): tl.add([a, b], 'right', 700 + i * 200, 500)
tl.add(16, 'rise', 1300, 500); tl.add(17, 'right', 1450, 400); tl.add([20, 23, 26, 29, 32], 'fade', 1500, 300)
for i, (bar, lab) in enumerate([(19, 18), (22, 21), (25, 24), (28, 27), (31, 30)]):
    tl.add(bar, 'up', 1550 + i * 130, 400); tl.add(lab, 'fade', 1850 + i * 130, 300)
tl.add(33, 'zoom', 2500, 500); tl.add(34, 'fade', 2800, 400)
tl.add(35, 'rise', 1500, 500)
rows = {}
for s in s3.shapes:
    if s.name.startswith('Oval'): rows.setdefault(round(s.top / 9525), []).append(s.shape_id)
for r, y in enumerate(sorted(rows)): tl.add(rows[y], 'zoom', 1800 + r * 70, 250)
tl.add(261, 'fade', 2850, 300); tl.add(262, 'zoom', 3050, 500); tl.add(263, 'fade', 3350, 400)
tl.add(264, 'rise', 1700, 500); tl.add(265, 'right', 2000, 400); tl.add([272, 273, 274, 275], 'fade', 2050, 300)
for i, (bar, lab) in enumerate([(267, 266), (269, 268), (271, 270)]):
    tl.add(bar, 'up', 2150 + i * 220, 450); tl.add(lab, 'fade', 2500 + i * 220, 300)
tl.add(276, 'zoom', 3500, 500); tl.add(277, 'fade', 3800, 400); tl.add([278, 279], 'fade', 3900, 400)
apply(s3, tl)

# 4장 추진 내용
tl = Tl(); header(tl)
for i, (a, b) in enumerate([(14, 16), (30, 32), (49, 51)]): tl.add([a, b], 'rise', 700 + i * 150, 500)
tl.add([17, 33, 52], 'fade', 1200, 300)
tl.add([84, 69], 'zoom', 1300, 500); tl.add([86, 67], 'zoom', 1450, 500); tl.add(71, 'zoom', 1600, 500)
tl.add(87, 'zoom', 1850, 350); tl.add(88, 'zoom', 1950, 350); tl.add(89, 'zoom', 2050, 350)
tl.add([75, 76, 77], 'zoom', 2250, 400); tl.add([78, 79, 80], 'zoom', 2350, 400); tl.add([81, 82, 83], 'zoom', 2450, 400)
tl.add([20, 24, 25, 26, 27], 'rise', 2600, 500); tl.add(66, 'rise', 2750, 500); tl.add(85, 'rise', 2900, 500)
tl.add(90, 'zoom', 3100, 350); tl.add(91, 'zoom', 3200, 350); tl.add(92, 'zoom', 3300, 350); tl.add(65, 'fade', 3300, 300)
apply(S(4), tl)

# 6장 기대효과
tl = Tl(); header(tl)
tl.add(10, 'rise', 700, 500); tl.add(23, 'rise', 850, 500); tl.add(35, 'rise', 1000, 500)
tl.add(11, 'fade', 900, 400); tl.add(24, 'fade', 1050, 400); tl.add(36, 'fade', 1200, 400)
tl.add(12, 'right', 1200, 400); tl.add([17, 18, 51], 'fade', 1300, 300)
for i, (bar, val) in enumerate([(14, 13), (16, 15), (49, 50)]):
    tl.add(bar, 'up', 1350 + i * 250, 450); tl.add(val, 'zoom', 1700 + i * 250, 350)
tl.add(19, 'fade', 2200, 300); tl.add(20, 'right', 2300, 400); tl.add(21, 'zoom', 2500, 500); tl.add(22, 'fade', 2850, 400)
tl.add(25, 'right', 1400, 400); tl.add([30, 31], 'fade', 1500, 300)
tl.add(27, 'up', 1550, 450); tl.add(26, 'zoom', 1900, 350); tl.add(29, 'up', 1850, 450); tl.add(28, 'zoom', 2200, 350)
tl.add(32, 'right', 2400, 400); tl.add(33, 'zoom', 2650, 500); tl.add(34, 'fade', 3000, 400)
tl.add(37, 'right', 1600, 400); tl.add([42, 43], 'fade', 1700, 300)
tl.add(39, 'up', 1750, 450); tl.add(38, 'zoom', 2100, 350); tl.add(41, 'right', 2100, 400); tl.add(40, 'zoom', 2400, 350)
tl.add(44, 'right', 2600, 400); tl.add(45, 'zoom', 2850, 500); tl.add(46, 'fade', 3200, 400)
tl.add([47, 48], 'fade', 3300, 400)
apply(S(6), tl)

# 7장 확산
tl = Tl(); header(tl)
tl.add([12, 65], 'rise', 700, 500); tl.add([16, 62], 'rise', 900, 500); tl.add([20, 64], 'rise', 1100, 500)
tl.add([10, 11], 'zoom', 1200, 400)
tl.add([14, 15], 'fade', 1000, 400); tl.add([18, 19], 'fade', 1200, 400); tl.add([22, 23], 'fade', 1400, 400)
tl.add(24, 'fade', 1600, 400); tl.add(25, 'right', 1700, 1500)
for i, (node, *labs) in enumerate([(26, 27, 28, 29), (30, 31, 32, 33), (34, 35, 36, 37), (38, 39, 40, 41), (42, 43, 44, 45)]):
    tl.add(node, 'zoom', 1750 + i * 300, 350); tl.add(list(labs), 'rise', 1850 + i * 300, 400)
tl.add(46, 'fade', 3300, 300)
apply(S(7), tl)

# 8장 마무리
tl = Tl()
tl.add(2, 'right', 0, 900); tl.add([3, 4, 5, 6], 'zoom', 200, 400); tl.add(7, 'rise', 400, 800)
tl.add(8, 'right', 1300, 500)
for i, ids in enumerate([[9, 10, 11], [12, 13, 14], [15, 16, 17]]): tl.add(ids, 'rise', 1700 + i * 200, 500)
tl.add(19, 'fade', 2500, 400); tl.add(18, 'rise', 2650, 500)
apply(S(8), tl)

for i, s in enumerate(p.slides, 1):
    transition(s, 'fade')
p.save(OUT); print('saved', OUT)
