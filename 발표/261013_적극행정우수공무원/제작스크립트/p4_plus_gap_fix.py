from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
P=lambda v:Emu(int(round(v*12192000/1920)));X=lambda e:e*1920/12192000
p=Presentation('fixed3.pptx');s=p.slides[3]
geo=lambda q:(X(q.left),X(q.top),X(q.width),X(q.height))
def setg(q,x,y,w,h):q.left,q.top,q.width,q.height=P(x),P(y),P(w),P(h)
def rm(q):q._element.getparent().remove(q._element)
def nostyle(q):
    st=q._element.find(qn('p:style'))
    if st is not None:q._element.remove(st)
cols=[110,690,1270];CW=540;H=256;R2=410;R3=734;GAPC=700;D=48
hl=[q for q in s.shapes if q.shape_type==1 and X(q.width)>500 and X(q.height)<80 and X(q.top)>400][0]
# pictures
for q in s.shapes:
    if q.shape_type!=13:continue
    x0,y0,w0,h0=geo(q);top=R2 if y0<700 else R3
    a0=w0/h0;a1=CW/H
    if a0<a1:
        keep=a0/a1;fh=1-q.crop_top-q.crop_bottom;d=fh*(1-keep)
        if q.name=='Picture 18': q.crop_top+=d; vy0=y0+h0*(1-keep)   # keep bottom alert bar
        else: q.crop_top+=d/2;q.crop_bottom+=d/2; vy0=y0+h0*(1-keep)/2
        vh=h0*keep
    else:
        keep=a1/a0;fw=1-q.crop_left-q.crop_right;d=fw*(1-keep);q.crop_left+=d/2;q.crop_right+=d/2;vy0=y0;vh=h0
    if q.name=='Picture 18':
        hx,hy,hw,hh=geo(hl);sy=H/vh;ny=top+(hy-vy0)*sy;nh=hh*sy
        ny=min(ny,top+H-4-nh);setg(hl,hx,ny,hw,nh)
    setg(q,x0,top,CW,H)
# placeholder + its texts
ph=[q for q in s.shapes if q.shape_type==1 and 'Rounded' in q.name and X(q.top)>700 and X(q.width)>500][0]
oy=X(ph.top);setg(ph,cols[0],R3,CW,H)
pts=sorted([q for q in s.shapes if q.has_text_frame and X(q.top)>700 and X(q.left)<650 and q.text_frame.text.strip()],key=lambda q:q.top)
tot=sum(X(q.height) for q in pts)+12*(len(pts)-1);yy=R3+(H-tot)/2
for q in pts:q.top=P(yy);yy+=X(q.height)+12
# replace + circles with circle + drawn cross (exact center)
for q in list(s.shapes):
    if q.shape_type==1 and q.has_text_frame and q.text_frame.text=='+':
        x,y,w,h=geo(q);cx=x+w/2;rm(q)
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,P(cx-D/2),P(GAPC-D/2),P(D),P(D));nostyle(c)
        c.fill.solid();c.fill.fore_color.rgb=RGBColor(0x2F,0x80,0xED);c.line.fill.background()
        L,T=24,5
        for (w_,h_) in ((L,T),(T,L)):
            b=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,P(cx-w_/2),P(GAPC-h_/2),P(w_),P(h_));nostyle(b)
            b.fill.solid();b.fill.fore_color.rgb=RGBColor(255,255,255);b.line.fill.background()
p.save('fixed4.pptx');print('ok')
# 정부24 화면은 위쪽(로고)을 살리고 아래를 자름
p=Presentation('fixed4.pptx');s=p.slides[3]
for q in s.shapes:
    if q.shape_type==13 and X(q.left)>1200 and X(q.top)>700:
        q.crop_bottom+=q.crop_top;q.crop_top=0.0
p.save('fixed4.pptx')
