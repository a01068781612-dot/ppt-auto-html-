from pptx import Presentation
from pptx.util import Emu,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.oxml.ns import qn
P=lambda v:Emu(int(round(v*12192000/1920)));X=lambda e:e*1920/12192000
p=Presentation('fixed2.pptx');s=p.slides[3]
geo=lambda q:(X(q.left),X(q.top),X(q.width),X(q.height))
def setg(q,x,y,w,h):q.left,q.top,q.width,q.height=P(x),P(y),P(w),P(h)
def rm(q):q._element.getparent().remove(q._element)
cols=[110,690,1270];CW=540;ROWS=[410,715];RH=275
# 1 title
t=[q for q in s.shapes if q.has_text_frame and '한 층' in q.text_frame.text and X(q.top)<240][0]
para=t.text_frame.paragraphs[0];runs=para.runs
runs[0].text='새로 만들지 않고, 기존 시스템에 한 층 더'
for r in runs[1:]:r._r.getparent().remove(r._r)
x,y,w,h=geo(t);t.width=P(1400)
# 3 remove frames in rows 2/3, labels beside placeholder
for q in list(s.shapes):
    x,y,w,h=geo(q)
    if q.shape_type==1 and 'Rounded' in q.name and w>500 and y>400 and h>200: rm(q)
    elif q.has_text_frame and q.text_frame.text.strip() in ('이미 있던 것','차량관제시스템'): rm(q)
# pictures -> full cells, same size; keep highlight following nav
hl=[q for q in s.shapes if q.shape_type==1 and X(q.width)>500 and X(q.height)<80 and X(q.top)>400][0]
def cell(x,y):
    c=min(range(3),key=lambda i:abs(x-(cols[i]+CW/2)));r=0 if y<700 else 1
    return cols[c],ROWS[r]
for q in s.shapes:
    if q.shape_type!=13:continue
    x0,y0,w0,h0=geo(q);cx,cy=cell(x0+w0/2,y0+h0/2)
    a0=w0/h0;a1=CW/RH;cl,ct,cr,cb=q.crop_left,q.crop_top,q.crop_right,q.crop_bottom
    fw,fh=1-cl-cr,1-ct-cb
    if a0>a1:
        keep=a1/a0;d=fw*(1-keep);q.crop_left=cl+d/2;q.crop_right=cr+d/2
        vx0=x0+w0*(1-keep)/2;vy0=y0;vw=w0*keep;vh=h0
    else:
        keep=a0/a1;d=fh*(1-keep);q.crop_top=ct+d/2;q.crop_bottom=cb+d/2
        vx0=x0;vy0=y0+h0*(1-keep)/2;vw=w0;vh=h0*keep
    if q.name=='Picture 18':
        hx,hy,hw,hh=geo(hl);sx=CW/vw;sy=RH/vh
        setg(hl,cx+(hx-vx0)*sx,cy+(hy-vy0)*sy,hw*sx,hh*sy)
    setg(q,cx,cy,CW,RH);q.line.fill.background()
# placeholder -> full cell, texts centered in it
ph=[q for q in s.shapes if q.shape_type==1 and 'Rounded' in q.name and X(q.width)<300 and X(q.top)>700][0]
setg(ph,cols[0],ROWS[1],CW,RH)
pts=[q for q in s.shapes if q.has_text_frame and X(q.top)>700 and X(q.left)<650 and q.text_frame.text.strip()]
pts.sort(key=lambda q:q.top);tot=sum(X(q.height) for q in pts)+12*(len(pts)-1);yy=ROWS[1]+(RH-tot)/2
for q in pts:
    h=X(q.height);setg(q,cols[0],yy,CW,h);yy+=h+12
    for para in q.text_frame.paragraphs:para.alignment=PP_ALIGN.CENTER
# 2 plus signs -> bold circle badges
for q in list(s.shapes):
    if q.has_text_frame and q.text_frame.text.strip()=='+':
        x,y,w,h=geo(q);cx=x+w/2;rm(q)
        D=52;c=s.shapes.add_shape(MSO_SHAPE.OVAL,P(cx-D/2),P(700-D/2),P(D),P(D))
        c.fill.solid();c.fill.fore_color.rgb=RGBColor(0x2F,0x80,0xED);c.line.color.rgb=RGBColor(255,255,255);c.line.width=Pt(3)
        st=c._element.find(qn('p:style'))
        if st is not None:c._element.remove(st)
        tf=c.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        pa=tf.paragraphs[0];pa.alignment=PP_ALIGN.CENTER;r=pa.add_run();r.text='+'
        f=r.font;f.size=Pt(26);f.bold=True;f.color.rgb=RGBColor(255,255,255);f.name='Pretendard'
# highlight: keep inside nav picture, under the + badges
hx,hy,hw,hh=geo(hl);bottom=min(hy+hh,ROWS[0]+RH-4);setg(hl,cols[0]+4,bottom-hh,CW-8,hh)
tree=s.shapes._spTree
el=hl._element;el.getparent().remove(el);tree.append(el)
for q in list(s.shapes):
    if q.shape_type==1 and q.has_text_frame and q.text_frame.text=='+':
        e=q._element;e.getparent().remove(e);tree.append(e)
p.save('fixed3.pptx');print('ok')
