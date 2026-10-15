# 1차 초안 v.2.1.1 → v.2.2.0: 대본(4:30) 확정 내용 반영
import copy, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC, OUT = sys.argv[1], sys.argv[2]
P = lambda v: Emu(int(round(v * 9525)))
p = Presentation(SRC)
S = lambda n: p.slides[n - 1]

def sh(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)

def set_text(s, text, run=0):
    rs = s.text_frame.paragraphs[0].runs
    rs[run].text = text

def dup(slide, s):
    el = copy.deepcopy(s._element)
    nid = max(int(x) for x in slide.shapes._spTree.xpath('//p:cNvPr/@id')) + 1
    el.xpath('.//p:cNvPr')[0].set('id', str(nid))
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]

def geom(s, x=None, y=None, w=None, h=None):
    if x is not None: s.left = P(x)
    if y is not None: s.top = P(y)
    if w is not None: s.width = P(w)
    if h is not None: s.height = P(h)

def tag(slide, x, y, text, fill):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, P(x), P(y), P(10), P(22))
    b.adjustments[0] = 0.5
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    b.line.fill.background(); b.shadow.inherit = False
    tf = b.text_frame
    tf.margin_left = tf.margin_right = P(9); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text
    f = r.font; f.size = Pt(9.5); f.color.rgb = RGBColor(255, 255, 255); f.name = 'Pretendard ExtraBold'
    rPr = r._r.get_or_add_rPr()
    for t in ('a:ea', 'a:cs'):
        e = rPr.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}' + t[2:], {'typeface': 'Pretendard ExtraBold'})
        rPr.append(e)
    # 글자 폭에 맞춰 너비 계산(한글 약 13px, 그 외 약 7px @9.5pt)
    w = sum(13 if '가' <= c <= '힣' else 7 for c in text) + 20
    b.width = P(w)
    return b

# ---------- 3장: 진입 불가 주체를 쉽게 ----------
s3 = S(3)
t = sh(s3, 261); set_text(t, '● 중형 소방차 진입 불가 164'); geom(t, w=330)

# ---------- 4장: 예산·외주 배지 삭제 + 기존/한 층 표시 ----------
s4 = S(4)
for sid in (10, 11, 12, 13):
    e = sh(s4, sid)._element; e.getparent().remove(e)
B, G = '2F80ED', '4B5563'
for x, y, txt in [(79, 271, '한 층 더 · 내비 알림'), (467, 271, '한 층 더 · AI 검색'), (855, 281, '한 층 더 · 재난 6개 기관 연결')]:
    tag(s4, x, y, txt, B)
for x, y, txt in [(80, 496, '기존 · 상황실 차량관제'), (468, 497, '기존 · 대원 제작 마을 책자'), (855, 496, '기존 · 정부24 원스톱')]:
    tag(s4, x, y, txt, G)

# ---------- 6장: 출동 그래프 3막대(미적용·우선신호·+내비 알림) ----------
s6 = S(6)
base = 433; H = 147
bars = [(sh(s6, 14), 100, 'C3CEDD'), (sh(s6, 16), 43.7, '9DC1F3'), (None, 40, '2F80ED')]
b3 = dup(s6, sh(s6, 16))
bars[2] = (b3, 40, '2F80ED')
vals = [sh(s6, 13), sh(s6, 15), dup(s6, sh(s6, 15))]
labs = [sh(s6, 17), sh(s6, 18), dup(s6, sh(s6, 18))]
xs = [121, 221, 321]; W = 72
for i, ((bar, v, col), vt, lb, x) in enumerate(zip(bars, vals, labs, xs)):
    h = H * v / 100
    geom(bar, x, base - h, W, h)
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(col)
    vh = 25 if i == 0 else 32
    geom(vt, x - 14, base - h - vh - 6, W + 28, vh)
    geom(lb, x - 14, 440, W + 28, 15)
set_text(vals[1], '43.7'); vals[1].text_frame.paragraphs[0].runs[0].font.size = Pt(16)
vals[1].text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('5B8FD8')
set_text(vals[2], '40')
set_text(labs[2], '+ 내비 알림')
labs[2].text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('2F80ED')
cap = sh(s6, 19); set_text(cap, '1km 통행시간 지수 · 우선신호 국내 실측 · 내비 알림 추정'); geom(cap, 93, 463, 328)
big = sh(s6, 21); set_text(big, '56% → 60%↓'); big.text_frame.paragraphs[0].runs[0].font.size = Pt(33)
sub = sh(s6, 22); set_text(sub, '5km 출동마다 30초 이상 단축 기대'); geom(sub, w=300)
src = sh(s6, 47)
full = ''.join(r.text for r in src.text_frame.paragraphs[0].runs)
i2 = full.index('②')
new = "출처: ① 우선신호 국내 실측(수원 56.3%·청주 55.4%·경기 61.3%↓) + 내비 알림 효과 추정(Weibull 외 2024: 알림 없으면 운전자 절반 가까이 미양보) " + full[i2:]
rs = src.text_frame.paragraphs[0].runs
rs[0].text = new
for r in rs[1:]: r.text = ''

# ---------- 7장: 원패스 → 개발 착수 ----------
s7 = S(7)
t = sh(s7, 23); set_text(t, '행정안전부장관상(대상) → 정부24 개발 착수'); geom(t, w=360)

# ---------- 8장: 영상2 엔딩과 같은 마무리 문장 ----------
s8 = S(8)
t = sh(s8, 7)
p2 = t.text_frame.paragraphs[1]
r1, r2 = p2.runs[0], p2.runs[1]
# r1(파랑) ↔ r2(검정) 순서 바꿔 "함께 지킬 때 "(검정) + "앞당겨집니다"(파랑)
r1.text, r2.text = '함께 지킬 때 ', '앞당겨집니다'
c1 = r1.font.color.rgb; r1.font.color.rgb = r2.font.color.rgb; r2.font.color.rgb = c1
ul = sh(s8, 8); geom(ul, 100 + 372, w=396)
for sid, txt in [(11, '길을 비켜 준 시민'), (14, '마을과 마주한 대원'), (17, '손 내민 정부')]:
    set_text(sh(s8, sid), txt)

# ---------- 발표자 노트: 확정 대본(남성·천천히, 4:30) ----------
NOTES = {
1: "[0:00–0:14]\n안녕하십니까, 경상남도 소방본부 119종합상황실 박중휘입니다. (0:08)\n재난 골든타임, 세 번 앞당기다. (0:12)\n먼저 영상입니다. (0:14)",
2: "[0:14–0:59] 영상1 · 45초 — 멘트 없음\n삽입 → 비디오 → 이 디바이스, 화면 전체 크기, 재생: 자동 실행.",
3: "[0:59–1:27]\n골든타임은 세 번 늦어집니다. (1:03)\n첫째, 출동. 막힌 길에서 신호를 어긴 긴급차량이 2.3배 늘었습니다. (1:11)\n둘째, 현장대응. 농촌 마을 164곳은 불을 끄는 중형 소방차가 못 들어갑니다. (1:20)\n셋째, 복구. 불이 꺼져도 6기관 8회, 90일을 기다립니다. (1:27)",
4: "[1:27–2:05]\n그래서 새로 만들지 않고, 기존 시스템에 한 층 더 얹었습니다. (1:34)\n첫째, 상황실 출동 정보를 카카오내비에 연결해, 긴급차량 접근 알림을 띄웁니다. (1:44)\n둘째, 대원들이 만든 마을 길 책자에 AI를 얹어, 5초 만에 답합니다. (1:52)\n셋째, 정부24 원스톱 서비스에 '재난'을 더해, 한 번 신청으로 6개 기관을 잇습니다. (2:02)\n실제 작동 화면입니다. (2:05)",
5: "[2:05–3:24] 영상2 · 1분 20초 — 멘트 없음 (엔딩 TTS: 골든타임은, 함께 지킬 때 앞당겨집니다.)\n삽입 → 비디오, 화면 전체 크기, 자동 실행.",
6: "[3:24–3:56]\n세 번의 골든타임, 이렇게 앞당깁니다. (3:29)\n첫째, 출동. 우선신호로 56% 줄어든 통행시간이, 내비 알림을 더하면 약 60%까지 줄어들 것으로 기대합니다. (3:43)\n둘째, 현장대응. 출동에서 방수까지 3분 단축입니다. (3:50)\n셋째, 복구. 14일이 7분이 되도록 설계했습니다. (3:56)\n\n[질의응답 근거] 1km 200초 → 우선신호 87초(56%). 길이 비었을 때 60초, 남은 27초의 절반(13.5초)을 앞차 미양보 지연으로 보고, 알림 받는 운전자 50%가 비키면 6.75초 단축 → 80초(약 60%). 5km 출동 시 약 34초. 30%면 58.5%, 80%(현대·기아·티맵 확대)면 62%.",
7: "[3:56–4:17]\n경남에서 먼저 증명했습니다. (3:59)\n내비 알림 497대, 농골길 18개 서 공유. (4:05)\n재난 원패스는 행정안전부장관상 대상을 받아, 정부24 개발에 착수했습니다. (4:14)\n이제 전국으로 갑니다. (4:17)",
8: "[4:17–4:30]\n길을 비켜 준 시민, 마을과 마주한 대원, 손 내민 정부. (4:23)\n골든타임은, 함께 지킬 때 앞당겨집니다. (4:28)\n감사합니다. (4:30)",
}
for n, txt in NOTES.items():
    S(n).notes_slide.notes_text_frame.text = txt

p.save(OUT)
print('saved', OUT)
